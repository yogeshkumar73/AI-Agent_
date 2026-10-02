import asyncio
import io
import docx
from pypdf import PdfWriter
from app.main import app
from app.core.database import connect_to_mongo, close_mongo_connection
from httpx import AsyncClient, ASGITransport

def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("QUARTERLY UTILITY BILL STATEMENT", level=1)
    doc.add_paragraph("Issued by: Apex Power & Electric Utilities Inc.")
    doc.add_paragraph("Invoice No: APEX-2026-9941")
    doc.add_paragraph("Date: 2026-09-30")
    doc.add_paragraph("Due Date: 2026-10-25")
    doc.add_paragraph("Total Amount Due: $4,850.75")
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Service Description"
    hdr_cells[1].text = "Units (kWh)"
    hdr_cells[2].text = "Charge ($)"
    
    row1 = table.add_row().cells
    row1[0].text = "Commercial Electricity Grid Consumption"
    row1[1].text = "28,400"
    row1[2].text = "$3,950.00"
    
    row2 = table.add_row().cells
    row2[0].text = "Renewable Energy Transition Surcharge"
    row2[1].text = "Standard"
    row2[2].text = "$900.75"

    doc.add_paragraph("Please remit payments prior to the due date to avoid a 5% late fee surcharge.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

async def run_e2e_test():
    print("\n--- Starting End-to-End System Test ---")
    await connect_to_mongo()
    transport = ASGITransport(app=app)
    
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Test Health
        health_res = await client.get("/health")
        print(f"1. Health Check: {health_res.status_code} -> {health_res.json()}")
        assert health_res.status_code == 200

        # 2. Register Client 1 (Acme Corp)
        reg1_res = await client.post("/api/v1/auth/register", json={
            "email": "client1@acmecorp.com",
            "password": "password123",
            "full_name": "Alice Henderson",
            "company_name": "Acme Corporation"
        })
        print(f"2. Register Client 1: {reg1_res.status_code}")
        token1 = reg1_res.json().get("access_token")
        headers1 = {"Authorization": f"Bearer {token1}"}

        # 3. Upload Sample DOCX Bill
        file_bytes = create_sample_docx()
        files = {"file": ("Apex_Electric_Bill.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        up_res = await client.post("/api/v1/documents/upload", files=files, headers=headers1)
        print(f"3. Upload Document: {up_res.status_code} -> ID: {up_res.json().get('id')}")
        doc_id = up_res.json()["id"]

        # Wait a moment for background processing to finish
        await asyncio.sleep(2.0)

        # 4. Check Document Status & Extracted KPIs
        doc_res = await client.get(f"/api/v1/documents/{doc_id}", headers=headers1)
        doc_data = doc_res.json()
        print(f"4. Parsed Status: {doc_data['status']}")
        print(f"   Extracted Entity: {doc_data['extracted_metadata'].get('entity_name')}")
        print(f"   Extracted Total: ${doc_data['extracted_metadata'].get('total_amount')}")
        print(f"   Invoice Ref: {doc_data['extracted_metadata'].get('invoice_number')}")
        print(f"   Summary: {doc_data.get('summary')}")

        # 5. Check Insights
        ins_res = await client.get("/api/v1/insights", headers=headers1)
        print(f"5. Generated AI Insights: {len(ins_res.json())} findings discovered.")
        for ins in ins_res.json():
            print(f"   * [{ins['insight_type']}] {ins['title']}: {ins['description']}")

        # 6. Test AI Agent Q&A with Citation
        conv_res = await client.post("/api/v1/chat/conversations", json={"title": "Utility Analysis", "document_id": doc_id}, headers=headers1)
        conv_id = conv_res.json()["id"]

        msg_res = await client.post(
            f"/api/v1/chat/conversations/{conv_id}/messages",
            json={"content": "What is the total amount due on this bill and what are the services?"},
            headers=headers1
        )
        msg_data = msg_res.json()
        print(f"\n6. AI Agent Response:")
        print(f"   Answer: {msg_data['content']}")
        print(f"   Citations: {msg_data['citations']}")
        print(f"   Token Usage: {msg_data['token_usage']}")
        print(f"   Follow-up suggestions: {msg_data['suggested_followups']}")

        # 7. Test Multi-User Isolation (Client 2 cannot see Client 1's documents)
        reg2_res = await client.post("/api/v1/auth/register", json={
            "email": "client2@globex.com",
            "password": "password123",
            "full_name": "Bob Martinez",
            "company_name": "Globex International"
        })
        token2 = reg2_res.json()["access_token"]
        headers2 = {"Authorization": f"Bearer {token2}"}

        # Client 2 tries to list documents -> should be 0!
        c2_docs = await client.get("/api/v1/documents", headers=headers2)
        print(f"\n7. Multi-Tenant Isolation Check:")
        print(f"   Client 2 sees {len(c2_docs.json())} documents (Expected: 0).")
        assert len(c2_docs.json()) == 0

        # Client 2 tries to access Client 1's document directly -> should be 404/denied
        c2_access_c1 = await client.get(f"/api/v1/documents/{doc_id}", headers=headers2)
        print(f"   Client 2 direct access to Client 1 doc returned status: {c2_access_c1.status_code} (Expected 404).")
        assert c2_access_c1.status_code == 404

    await close_mongo_connection()
    print("\n--- ALL END-TO-END TESTS PASSED SUCCESSFULLY! ---\n")

if __name__ == "__main__":
    asyncio.run(run_e2e_test())
